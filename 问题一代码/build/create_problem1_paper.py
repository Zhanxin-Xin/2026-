from __future__ import annotations

import csv
import math
from pathlib import Path

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "build"
EQ = BUILD / "equations"
DRAFT = BUILD / "problem1_draft.docx"
FLOW = BUILD / "problem1_pipeline.png"
SUMMARY = ROOT / "outputs" / "summary.csv"
TYPICAL = ROOT / "outputs" / "typical_alignment_figure.png"
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")

MATHML = {
    "eq1": r'''<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msub><mi>I</mi><mi>k</mi></msub><mo>=</mo><mo>[</mo><mi>k</mi><mi>Δt</mi><mo>,</mo><mi>min</mi><mo>(</mo><mo>(</mo><mi>k</mi><mo>+</mo><mn>1</mn><mo>)</mo><mi>Δt</mi><mo>,</mo><mi>T</mi><mo>)</mo><mo>)</mo><mo>,</mo><mspace width="1em"/><mi>N</mi><mo>=</mo><mi>min</mi><mo>(</mo><mn>70</mn><mo>,</mo><mo>⌈</mo><mi>T</mi><mo>/</mo><mi>Δt</mi><mo>⌉</mo><mo>)</mo><mo>,</mo><mspace width="1em"/><mi>Δt</mi><mo>=</mo><mn>0.5</mn><mspace width="0.3em"/><mi mathvariant="normal">s</mi><mo>.</mo></mrow></math>''',
    "eq2": r'''<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msub><mi mathvariant="bold">e</mi><mi>j</mi></msub><mo>=</mo><mfrac><mn>1</mn><mrow><mo>|</mo><msub><mi>S</mi><mi>j</mi></msub><mo>|</mo></mrow></mfrac><munder><mo>∑</mo><mrow><mi>r</mi><mo>∈</mo><msub><mi>S</mi><mi>j</mi></msub></mrow></munder><msub><mi mathvariant="bold">h</mi><mi>r</mi></msub><mo>,</mo><mspace width="1em"/><msub><mi mathvariant="bold">e</mi><mi>j</mi></msub><mo>∈</mo><msup><mi>R</mi><mn>768</mn></msup><mo>.</mo></mrow></math>''',
    "eq3": r'''<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msubsup><mi mathvariant="bold">x</mi><mi>k</mi><mo>(T)</mo></msubsup><mo>=</mo><mfrac><mrow><munder><mo>∑</mo><mi>j</mi></munder><msub><mi>ℓ</mi><mrow><mi>j</mi><mi>k</mi></mrow></msub><msub><mi mathvariant="bold">e</mi><mi>j</mi></msub></mrow><mrow><munder><mo>∑</mo><mi>j</mi></munder><msub><mi>ℓ</mi><mrow><mi>j</mi><mi>k</mi></mrow></msub></mrow></mfrac><mo>,</mo><mspace width="1em"/><msub><mi>ℓ</mi><mrow><mi>j</mi><mi>k</mi></mrow></msub><mo>=</mo><mo>|</mo><mo>[</mo><msub><mi>a</mi><mi>j</mi></msub><mo>,</mo><msub><mi>b</mi><mi>j</mi></msub><mo>)</mo><mo>∩</mo><msub><mi>I</mi><mi>k</mi></msub><mo>|</mo><mo>.</mo></mrow></math>''',
    "eq4": r'''<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msubsup><mi mathvariant="bold">x</mi><mi>k</mi><mo>(A)</mo></msubsup><mo>=</mo><mo>[</mo><msub><mi mathvariant="bold">μ</mi><mi>k</mi></msub><mo>;</mo><msub><mi mathvariant="bold">σ</mi><mi>k</mi></msub><mo>]</mo><mo>∈</mo><msup><mi>R</mi><mn>50</mn></msup><mo>,</mo><mspace width="0.8em"/><msub><mi mathvariant="bold">μ</mi><mi>k</mi></msub><mo>=</mo><msub><mi mathvariant="normal">mean</mi><mrow><mi>q</mi><mo>∈</mo><msub><mi>I</mi><mi>k</mi></msub></mrow></msub><msub><mi mathvariant="bold">z</mi><mi>q</mi></msub><mo>,</mo><mspace width="0.8em"/><msub><mi mathvariant="bold">σ</mi><mi>k</mi></msub><mo>=</mo><msub><mi mathvariant="normal">std</mi><mrow><mi>q</mi><mo>∈</mo><msub><mi>I</mi><mi>k</mi></msub></mrow></msub><msub><mi mathvariant="bold">z</mi><mi>q</mi></msub><mo>.</mo></mrow></math>''',
    "eq5": r'''<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msubsup><mi>m</mi><mi>k</mi><mo>(V)</mo></msubsup><mo>=</mo><mi>𝕀</mi><mo>(</mo><msub><mi>s</mi><mi>k</mi></msub><mo>=</mo><mn>1</mn><mo>)</mo><mi>𝕀</mi><mo>(</mo><msub><mi>c</mi><mi>k</mi></msub><mo>≥</mo><mn>0.8</mn><mo>)</mo><mi>𝕀</mi><mo>(</mo><msubsup><mi mathvariant="bold">x</mi><mi>k</mi><mo>(V)</mo></msubsup><mspace width="0.3em"/><mi mathvariant="normal">is finite</mi><mo>)</mo><mo>.</mo></mrow></math>''',
    "eq6": r'''<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msup><mi mathvariant="bold">X</mi><mo>(T)</mo></msup><mo>∈</mo><msup><mi>R</mi><mrow><mn>70</mn><mo>×</mo><mn>768</mn></mrow></msup><mo>,</mo><mspace width="1em"/><msup><mi mathvariant="bold">X</mi><mo>(A)</mo></msup><mo>∈</mo><msup><mi>R</mi><mrow><mn>70</mn><mo>×</mo><mn>50</mn></mrow></msup><mo>,</mo><mspace width="1em"/><msup><mi mathvariant="bold">X</mi><mo>(V)</mo></msup><mo>∈</mo><msup><mi>R</mi><mrow><mn>70</mn><mo>×</mo><mn>465</mn></mrow></msup><mo>.</mo></mrow></math>''',
    "eq7": r'''<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msubsup><mover><mi mathvariant="bold">x</mi><mo>~</mo></mover><mi>k</mi><mo>(m)</mo></msubsup><mo>=</mo><msubsup><mi>m</mi><mi>k</mi><mo>(m)</mo></msubsup><msubsup><mi>m</mi><mi>k</mi><mo>(P)</mo></msubsup><msubsup><mi mathvariant="bold">x</mi><mi>k</mi><mo>(m)</mo></msubsup><mo>,</mo><mspace width="1em"/><mi>m</mi><mo>∈</mo><mo>{</mo><mi>T</mi><mo>,</mo><mi>A</mi><mo>,</mo><mi>V</mi><mo>}</mo><mo>.</mo></mrow></math>''',
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, *, bold=False, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = "Times New Roman"
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_table_borders(table, color="BFBFBF", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_layout_fixed(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_body(doc, text, *, first_line=True, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        r1.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    keep_with_next(p)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_equation(doc, filename, number, width_cm):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(14.6)
    table.columns[1].width = Cm(1.0)
    table.cell(0, 0).width = Cm(14.6)
    table.cell(0, 1).width = Cm(1.0)
    for cell in table.row_cells(0):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=15, bottom=15, start=10, end=10)
    p = table.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    key = Path(filename).stem
    transform = etree.XSLT(etree.parse(str(MML2OMML)))
    omml = transform(etree.fromstring(MATHML[key].encode("utf-8"))).getroot()
    p._p.append(omml)
    p2 = table.cell(0, 1).paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run(number)
    r.font.size = Pt(10.5)
    r.font.name = "Times New Roman"
    # No visible borders for equation layout.
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)
    return table


def make_flowchart():
    width, height = 2200, 1160
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    font = ImageFont.truetype(str(font_path), 38)
    small = ImageFont.truetype(str(font_path), 34)
    boxes = [
        (40, 70, 390, 210, "标签表与\n100条MP4", "#E9EEF6"),
        (560, 70, 390, 210, "媒体探测与\n音轨解码", "#E9EEF6"),
        (1080, 70, 430, 210, "统一0.5秒\n原视频时间轴", "#DCE9F7"),
        (80, 450, 530, 250, "文本分支\nWhisperX强制对齐\nRoBERTa 768维", "#EDF4FA"),
        (835, 450, 530, 250, "语音分支\neGeMAPS 25 LLD\n均值+标准差 50维", "#E8F4F2"),
        (1590, 450, 530, 250, "视觉分支\n2 fps + OpenFace\n465维", "#F7ECEB"),
        (420, 850, 630, 230, "窗口聚合、有效性掩码\n零填充至70窗", "#E9EEF6"),
        (1320, 850, 620, 230, "NPZ特征 + 元数据\n日志、时间线与QC", "#E9EEF6"),
    ]
    for x, y, w, h, txt, color in boxes:
        draw.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=color, outline="#4C566A", width=4)
        bbox = draw.multiline_textbbox((0, 0), txt, font=font, spacing=10, align="center")
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.multiline_text((x + (w - tw) / 2, y + (h - th) / 2), txt, font=font, fill="#111111", spacing=10, align="center")

    def arrow(x1, y1, x2, y2):
        draw.line((x1, y1, x2, y2), fill="#4C566A", width=5)
        angle = math.atan2(y2 - y1, x2 - x1)
        tip = (x2, y2)
        left = (x2 - 24 * math.cos(angle - 0.55), y2 - 24 * math.sin(angle - 0.55))
        right = (x2 - 24 * math.cos(angle + 0.55), y2 - 24 * math.sin(angle + 0.55))
        draw.polygon([tip, left, right], fill="#4C566A")

    arrow(430, 175, 560, 175)
    arrow(950, 175, 1080, 175)
    for tx in (345, 1100, 1855):
        arrow(1295, 280, tx, 450)
    for sx in (345, 1100, 1855):
        arrow(sx, 700, 735, 850)
    arrow(1050, 965, 1320, 965)
    draw.text((1560, 155), "统一主键：video_id + clip_id", font=small, fill="#333333")
    img.save(FLOW, dpi=(300, 300))


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.45)
    section.bottom_margin = Cm(2.25)
    section.left_margin = Cm(2.55)
    section.right_margin = Cm(2.35)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for level, size, font_cn, before, after in [
        (1, 14, "黑体", 14, 7),
        (2, 12, "黑体", 10, 5),
        (3, 11, "黑体", 7, 3),
    ]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_cn)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = False

    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(9.5)
    caption.font.bold = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(4)


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)


def add_summary_stats_table(doc, df):
    add_caption(doc, "表 3-3  问题一全量特征提取与质量统计")
    rows = [
        ("样本覆盖", "100/100", "标签行与视频文件一一匹配，重复键与缺失视频均为0"),
        ("视频时长", f"{df.duration_seconds.min():.3f}--{df.duration_seconds.max():.3f} s", f"平均 {df.duration_seconds.mean():.3f} s，中位数 {df.duration_seconds.median():.3f} s"),
        ("统一序列", "0.5 s/窗，最多70窗", f"真实时间窗合计 {int(df.valid_length.sum())} 个"),
        ("文本特征", "70×768", f"有效窗 {int(df.text_valid_windows.sum())} 个；有效对齐词 1920/1932"),
        ("语音特征", "70×50", f"有效非静音窗 {int(df.audio_valid_windows.sum())} 个"),
        ("视觉特征", "70×465", "有效窗 1411 个；通过阈值帧 1411/1570"),
        ("质量核验", "全部通过", "形状、有限值、掩码、填充、时间越界与追溯文件检查均通过"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(3.0), Cm(3.8), Cm(9.0)]
    headers = ["统计项目", "结果", "说明"]
    for j, h in enumerate(headers):
        table.columns[j].width = widths[j]
        set_cell_text(table.cell(0, j), h, bold=True, size=9.5, color="FFFFFF")
        set_cell_shading(table.cell(0, j), "365F91")
    set_repeat_table_header(table.rows[0])
    for i, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for j, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if j == 2 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cells[j], val, size=9.2, align=align)
            cells[j].width = widths[j]
            if i % 2 == 0:
                set_cell_shading(cells[j], "F2F6FA")
    set_table_borders(table, color="D9D9D9", size="5")
    set_table_layout_fixed(table)


def add_file_table(doc):
    add_caption(doc, "表 3-1  自生成特征文件的组织结构与用途")
    rows = [
        ("features.npz", "三模态矩阵、四类掩码、窗口边界、有效长度及样本主键", "模型读取的主文件"),
        ("metadata.json", "媒体信息、特征名、模型与聚合规则、只读标签", "解释维度与参数"),
        ("word_alignment.csv", "词序号、词文本、起止时间、置信度与有效标志", "文本回溯"),
        ("timeline.json", "统一时间轴上的词、采样帧、人脸置信度与静音窗", "跨模态核验"),
        ("status.json", "处理签名、分阶段状态、告警、错误和工具版本", "断点续跑与审计"),
        ("summary.csv", "100条样本的宽表汇总", "论文统计"),
        ("modality_summary.csv", "样本×模态的300行长表", "全量逐模态清单"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(3.4), Cm(8.2), Cm(4.2)]
    for j, h in enumerate(["文件", "主要字段", "核验用途"]):
        table.columns[j].width = widths[j]
        set_cell_text(table.cell(0, j), h, bold=True, size=9.5, color="FFFFFF")
        set_cell_shading(table.cell(0, j), "365F91")
    set_repeat_table_header(table.rows[0])
    for i, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for j, val in enumerate(row):
            set_cell_text(cells[j], val, size=9.0, align=WD_ALIGN_PARAGRAPH.LEFT if j != 2 else WD_ALIGN_PARAGRAPH.CENTER)
            cells[j].width = widths[j]
            if i % 2 == 0:
                set_cell_shading(cells[j], "F2F6FA")
    set_table_borders(table, color="D9D9D9", size="5")
    set_table_layout_fixed(table)


def add_typical_table(doc):
    rows = [
        ("样本主键", "-dxfTGcXJoc__6"),
        ("原始时长与序列", "12.967 s，26个真实窗口"),
        ("文本", "27/27个词得到有效时间戳，25个窗口有效"),
        ("语音", "26个窗口有效，图中展示F0半音均值"),
        ("视觉", "25个窗口有效，图中展示AU12_r"),
        ("综合覆盖", "三模态窗口覆盖率均值97.4%"),
    ]
    add_caption(doc, "表 3-2  典型样本的对齐核验结果")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(4.0), Cm(11.8)]
    for j, h in enumerate(["项目", "核验结果"]):
        table.columns[j].width = widths[j]
        set_cell_text(table.cell(0, j), h, bold=True, size=9.5, color="FFFFFF")
        set_cell_shading(table.cell(0, j), "365F91")
    for i, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for j, val in enumerate(row):
            set_cell_text(cells[j], val, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT if j == 1 else WD_ALIGN_PARAGRAPH.CENTER)
            cells[j].width = widths[j]
            if i % 2 == 0:
                set_cell_shading(cells[j], "F2F6FA")
    set_table_borders(table, color="D9D9D9", size="5")
    set_table_layout_fixed(table)


def add_all_samples_table(doc, df):
    add_caption(doc, "表 3-4  附件1全部100条样本的特征提取结果")
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(5)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("注：三模态特征维度固定为T/A/V=768/50/465，对齐粒度固定为0.5 s；N为真实时间窗数，T、A、V分别为三模态有效窗数。")
    r.font.size = Pt(9)

    headers = ["序号", "样本编号", "时长/s", "N", "T", "A", "V", "状态"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(1.0), Cm(5.0), Cm(1.8), Cm(1.25), Cm(1.25), Cm(1.25), Cm(1.25), Cm(2.0)]
    for j, h in enumerate(headers):
        table.columns[j].width = widths[j]
        set_cell_text(table.cell(0, j), h, bold=True, size=8.6, color="FFFFFF")
        set_cell_shading(table.cell(0, j), "365F91")
    set_repeat_table_header(table.rows[0])
    for i, row in df.reset_index(drop=True).iterrows():
        vals = [
            i + 1,
            f"{row.video_id}__{str(row.clip_id)}",
            f"{row.duration_seconds:.3f}",
            int(row.valid_length),
            int(row.text_valid_windows),
            int(row.audio_valid_windows),
            int(row.visual_valid_windows),
            "成功" if row.status == "success" else str(row.status),
        ]
        cells = table.add_row().cells
        table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for j, val in enumerate(vals):
            set_cell_text(cells[j], val, size=8.0, align=WD_ALIGN_PARAGRAPH.LEFT if j == 1 else WD_ALIGN_PARAGRAPH.CENTER)
            cells[j].width = widths[j]
            if i % 2 == 1:
                set_cell_shading(cells[j], "F5F8FB")
    set_table_borders(table, color="D9D9D9", size="4")
    set_table_layout_fixed(table)


def add_versions_table(doc):
    rows = [
        ("运行环境", "Ubuntu 22.04；Python 3.8.20；随机种子2026"),
        ("媒体处理", "FFmpeg/FFprobe 4.4.2；音频16 kHz单声道PCM"),
        ("文本", "WhisperX 3.2.0；RoBERTa-base；Transformers 4.39.3；PyTorch 2.4.1"),
        ("语音", "openSMILE 2.5.0；eGeMAPSv02 LowLevelDescriptors"),
        ("视觉", "OpenFace 2.2.0；2 fps；人脸置信度阈值0.80"),
        ("数据依赖", "NumPy 1.24.3；pandas 2.0.3；openpyxl 3.1.5；SoundFile 0.12.1"),
    ]
    add_caption(doc, "表 3-5  问题一复现实验环境与关键参数")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(3.5), Cm(12.3)]
    for j, h in enumerate(["类别", "版本与参数"]):
        table.columns[j].width = widths[j]
        set_cell_text(table.cell(0, j), h, bold=True, size=9.5, color="FFFFFF")
        set_cell_shading(table.cell(0, j), "365F91")
    for i, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for j, val in enumerate(row):
            set_cell_text(cells[j], val, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT if j == 1 else WD_ALIGN_PARAGRAPH.CENTER)
            cells[j].width = widths[j]
            if i % 2 == 0:
                set_cell_shading(cells[j], "F2F6FA")
    set_table_borders(table, color="D9D9D9", size="5")
    set_table_layout_fixed(table)


def build():
    make_flowchart()
    df = pd.read_csv(SUMMARY)
    df["clip_id"] = df["clip_id"].astype(str)

    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("3.1 问题一的建模分析与求解")

    add_heading(doc, "3.1.1 问题描述与总体思路", 2)
    add_body(doc, "问题一要求由100条英文原始视频构造可直接复核的文本、语音和视觉时序特征。困难不只在于三类模态的特征尺度不同，还在于文本由离散词语组成、语音是高频连续波形、视频是图像序列，三者缺少天然一致的采样位置。因此，若先独立抽取定长向量再直接拼接，将无法回答某一特征行究竟对应原视频的哪个时段，也难以区分真实观测、局部无效和序列填充。")
    add_body(doc, "为满足完整性、可核验性和可复现性要求，本文建立以原视频绝对时间为主轴的可追溯三模态特征流水线。样本以video_id与clip_id组成的二元键唯一标识；标签表只用于建立清单和核对样本，不参与特征计算。每条视频先经过媒体探测与音轨解码，再在0.5 s统一时间窗上分别聚合词级语义、短时声学描述子和2 fps视觉特征，最后将全部序列零填充至70窗，并同时保存填充掩码、三模态有效性掩码、窗口边界和原始素材路径。总体流程如图3-1所示。")
    add_caption(doc, "图 3-1  问题一多模态特征提取与时序对齐流程")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FLOW), width=Cm(15.8))

    add_heading(doc, "3.1.2 统一时间轴与符号约定", 2)
    add_body(doc, "设一条样本的容器时长为T。考虑到题目中最长视频为34.567 s，本文取时间步长Δt=0.5 s，并将序列上限设为70，从而覆盖35 s。第k个时间窗及真实窗口数定义为")
    add_equation(doc, "eq1.png", "(3-1)", 12.8)
    add_body(doc, "其中k=0,1,…,N-1。所有模态均以I_k为唯一时间索引。若媒体流报告的结束时间与容器时长不一致，则使用容器时长建立主时间轴，同时保留各流报告时长和实际可解码帧数用于审计；超出可解码内容的视觉位置保持无效，不补造观测。由此，任何序列位置都能通过window_start与window_end还原到原视频的绝对时间区间。")

    add_heading(doc, "3.1.3 三模态情感特征构建", 2)
    add_heading(doc, "3.1.3.1 文本语义特征", 3)
    add_body(doc, "赛题给出的英文转写是唯一正式文本。本文不重新进行自动语音识别，而是利用WhisperX对赛题转写和原音轨做强制对齐，得到第j个词的起止时间[a_j,b_j)及置信分。无法匹配的词仍保留在记录中，但不给出虚构时间。随后使用RoBERTa-base最后一层隐状态提取上下文语义；同一词可能被分为多个子词，其词向量取对应子词向量的均值：")
    add_equation(doc, "eq2.png", "(3-2)", 7.5)
    add_body(doc, "其中S_j为第j个词对应的子词集合，h_r为第r个子词的最后一层隐向量。对落入时间窗I_k的词，按词区间与窗口的重叠长度加权聚合：")
    add_equation(doc, "eq3.png", "(3-3)", 8.8)
    add_body(doc, "当分母为0时，该窗文本向量置零且text_mask[k]=false；否则text_mask[k]=true。该规则允许跨窗词按实际覆盖时长分配语义信息，避免将整个词武断地归入单一窗口。")

    add_heading(doc, "3.1.3.2 语音韵律特征", 3)
    add_body(doc, "原视频音轨经FFmpeg解码为16 kHz、单声道PCM。采用openSMILE的eGeMAPSv02 LowLevelDescriptors提取25个短时低层描述子，包括响度、基频、谱通量、MFCC、抖动、闪烁、谐噪比和前三个共振峰等。对每个0.5 s窗口内的短时帧分别计算均值与标准差并串联：")
    add_equation(doc, "eq4.png", "(3-4)", 12.0)
    add_body(doc, "因此语音特征维度为50。若窗口内无可用短时帧、含非有限值，或音频RMS低于0.001，则将该位置记为静音/无效，audio_mask[k]=false；该窗仍保留在统一时间轴上，不删除样本。")

    add_heading(doc, "3.1.3.3 视觉表情与姿态特征", 3)
    add_body(doc, "视觉分支按2 fps从原视频确定性采样，即每个0.5 s窗口对应一个采样时刻。OpenFace 2.2.0输出视线方向与视角8维、眼部二维/三维关键点280维、头部平移与旋转6维、68点面部二维坐标136维，以及面部动作单元强度与出现状态35维，共465维。仅当OpenFace检测成功、置信度不低于0.80且全部特征为有限值时，视觉窗口才被视为有效：")
    add_equation(doc, "eq5.png", "(3-5)", 10.8)
    add_body(doc, "检测失败或低置信度只表示该时刻没有可信视觉观测，不解释为‘没有情绪’。对应特征置零并令visual_mask[k]=false，从而为后续鲁棒建模保留真实的局部缺失信息。")

    add_heading(doc, "3.1.4 跨模态时序对齐与定长组织", 2)
    add_body(doc, "三条分支最终都映射到同一组I_k。文本使用词区间与I_k的重叠长度，语音使用短时帧起点所属窗口，视觉使用2 fps采样序号与窗口序号的一一对应关系。对文本、语音和视觉任一模态，先计算该模态的有效掩码，再与真实时间轴掩码共同约束特征：")
    add_equation(doc, "eq7.png", "(3-6)", 8.4)
    add_body(doc, "其中真实时间轴掩码取1表示k<N，反之为填充位置。最终三模态矩阵的形状统一为")
    add_equation(doc, "eq6.png", "(3-7)", 12.4)
    add_body(doc, "固定形状便于批量建模，但零值本身不能判断该位置是否有效，读取时必须联合padding_mask与相应的text_mask、audio_mask或visual_mask。该双层掩码机制同时区分了‘视频已结束’与‘视频仍在继续但当前模态不可用’两种情况。")

    add_heading(doc, "3.1.5 特征文件规范与追溯设计", 2)
    add_body(doc, "每条样本单独存放于outputs/samples/<video_id>__<clip_id>/，主特征采用压缩NPZ格式，浮点特征统一为float32，掩码统一为布尔类型。NPZ中同时写入样本主键、源标签行号、原始时长、真实窗口数、70个窗口的起止时间和四类掩码。配套JSON与CSV保存模型、特征名、媒体属性、词级时间戳、帧级检测状态和阶段日志，形成‘特征行—时间窗—原始词/帧—源视频’的追溯链。各文件作用见表3-1。")
    add_file_table(doc)

    add_heading(doc, "3.1.6 典型样本的时序对齐验证", 2)
    add_body(doc, "从成功样本中自动选择三模态覆盖完整、词级对齐质量较高的-dxfTGcXJoc__6作为典型样本。该视频时长12.967 s，对应26个真实窗口；27个转写词全部获得有效起止时间。图3-2在同一原视频横轴上依次展示4个实际视频帧、词级区间、原始音频波形、窗口级F0半音均值、OpenFace动作单元AU12_r以及三模态覆盖掩码。浅蓝色区域给出一个三模态同时有效的0.5 s窗口，灰色斜线表示无效或越出真实时长的位置。")
    add_caption(doc, "图 3-2  典型样本的文本 语音 视觉时序对应关系")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(TYPICAL), width=Cm(16.0))
    add_typical_table(doc)
    add_body(doc, "图中词区间与音频波形的发声片段相互对应，视频帧时刻与视觉序列位置严格遵循2 fps规则；三模态均可由窗口索引还原到原素材。该样本的文本、语音和视觉有效窗分别为25、26和25，平均覆盖率为97.4%，说明统一时间轴没有改变各模态的原始时序关系。")

    add_heading(doc, "3.1.7 全量提取结果与质量核验", 2)
    add_body(doc, "程序对标签表与视频目录进行双向核对，100行标签均找到唯一视频，未发现重复主键、缺失视频、额外无标签视频、空文本或越界情感强度。全部100条样本均完成媒体探测、文本对齐、语音提取、视觉提取和文件写入，状态均为success。主要统计结果见表3-3。")
    add_summary_stats_table(doc, df)
    add_body(doc, "在1 932个对齐词中，1 920个具有有效时间戳，对齐有效率为99.38%；其余12个词分布在9条样本中，均保留原词并将相应位置记为无时间戳。视觉分支实际获得1 570个采样帧，其中1 411个通过检测成功与0.80置信度阈值，有效率为89.87%。所有词时间和帧时间的越界数均为0。按统一时间轴统计，100条样本共有2 287个真实窗口，文本、语音和视觉分别有1 814、1 538和1 411个有效窗口。视觉流短于容器时长或人脸不可见所造成的空窗均由掩码记录，未删除任何原始样本。")
    add_all_samples_table(doc, df)

    add_heading(doc, "3.1.8 可复现性与异常处理规则", 2)
    add_body(doc, "关键参数集中写入config.yaml，执行顺序固定为清单扫描、环境预检、单样本试跑、全量运行、质量核验和典型样本绘图。程序以配置、视频大小与修改时间、转写文本及流水线版本计算处理签名；签名不变且状态成功的样本可安全断点续跑，参数改变则自动重新处理。随机种子固定为2026。工具版本与关键参数见表3-5。")
    add_versions_table(doc)
    add_body(doc, "异常处理遵循‘保留原样本、显式标记、禁止替代’原则：词无法对齐时不伪造时间；静音、非有限声学帧和低置信度人脸均只关闭相应模态掩码；视频流时长冲突时保留容器时长、流结束时间和实际解码数量；工具或固定模型不可用时记录失败并继续处理其他样本，不以语义不同的特征替代。质量控制程序逐样本检查矩阵形状、数值有限性、零填充、掩码逻辑、时间边界、主键一致性和追溯文件完整性，最终报告structural_checks_passed=true、extraction_complete=true。")

    add_heading(doc, "3.1.9 问题一小结", 2)
    add_body(doc, "本文完成了附件1全部100条视频的三模态情感特征构建。以0.5 s原视频时间窗为公共坐标，将RoBERTa文本语义、eGeMAPS声学统计和OpenFace视觉行为统一组织为70步定长序列，并用填充掩码与模态掩码保存真实有效性。全量样本、全部模态文件和窗口位置均可追溯至原视频，典型样本验证与质量报告共同表明该特征集满足覆盖完整、时序可核验和过程可复现三项要求，可作为问题二和问题三的统一输入接口。")

    add_heading(doc, "问题一相关参考文献", 2)
    refs = [
        "[1] Liu Y, Ott M, Goyal N, et al. RoBERTa: A Robustly Optimized BERT Pretraining Approach. arXiv:1907.11692, 2019.",
        "[2] Bain M, Huh J, Han T, et al. WhisperX: Time-Accurate Speech Transcription of Long-Form Audio. INTERSPEECH, 2023: 4489-4493.",
        "[3] Eyben F, Wöllmer M, Schuller B. openSMILE: The Munich Versatile and Fast Open-Source Audio Feature Extractor. ACM Multimedia, 2010: 1459-1462.",
        "[4] Eyben F, Scherer K R, Schuller B W, et al. The Geneva Minimalistic Acoustic Parameter Set (GeMAPS) for Voice Research and Affective Computing. IEEE Transactions on Affective Computing, 2016, 7(2): 190-202.",
        "[5] Baltrušaitis T, Zadeh A, Lim Y C, et al. OpenFace 2.0: Facial Behavior Analysis Toolkit. IEEE International Conference on Automatic Face and Gesture Recognition, 2018: 59-66.",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(2)
        p.add_run(ref)

    # Apply page numbering to all sections and remove identifying metadata.
    for section in doc.sections:
        add_page_number(section)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "问题一 多模态情感特征提取与时序对齐"
    doc.core_properties.subject = "数学建模竞赛论文问题一"
    doc.core_properties.keywords = "多模态情感分析, 时序对齐, 特征提取"
    doc.save(DRAFT)


if __name__ == "__main__":
    build()
